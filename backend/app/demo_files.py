"""Build small PDF and DOCX files from synthetic policy text."""

from __future__ import annotations

import io
import re

from docx import Document


def build_text_pdf(text: str) -> bytes:
    lines = _wrap_lines(text)
    page_size = 42
    pages = [lines[index : index + page_size] for index in range(0, len(lines), page_size)] or [[]]
    objects: dict[int, bytes] = {
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    page_ids: list[int] = []
    next_id = 4
    for page_lines in pages:
        content = _content_stream(page_lines)
        content_id = next_id
        page_id = next_id + 1
        next_id += 2
        objects[content_id] = (
            b"<< /Length "
            + str(len(content)).encode("ascii")
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        )
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode("ascii")
        page_ids.append(page_id)
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[2] = f"<< /Type /Pages /Count {len(page_ids)} /Kids [{kids}] >>".encode("ascii")
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    return _serialize_pdf(objects)


def build_docx(text: str) -> bytes:
    document = Document()
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if re.match(r"^\d+\.\s+\S", line):
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _wrap_lines(text: str, width: int = 90) -> list[str]:
    lines: list[str] = []
    for paragraph in text.splitlines():
        if not paragraph.strip():
            lines.append("")
            continue
        current = ""
        for word in paragraph.split(" "):
            trial = word if not current else f"{current} {word}"
            if len(trial) > width and current:
                lines.append(current)
                current = word
            else:
                current = trial
        if current:
            lines.append(current)
    return lines


def _content_stream(lines: list[str]) -> bytes:
    commands = ["BT", "/F1 11 Tf", "50 760 Td", "14 TL"]
    for line in lines:
        commands.append(f"({_escape_pdf(line)}) Tj")
        commands.append("T*")
    commands.append("ET")
    return "\n".join(commands).encode("latin-1", errors="replace")


def _escape_pdf(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _serialize_pdf(objects: dict[int, bytes]) -> bytes:
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    max_id = max(objects)
    for object_id in range(1, max_id + 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("ascii"))
        output.extend(objects[object_id])
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {max_id + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for object_id in range(1, max_id + 1):
        output.extend(f"{offsets[object_id]:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n").encode(
            "ascii"
        )
    )
    return bytes(output)
